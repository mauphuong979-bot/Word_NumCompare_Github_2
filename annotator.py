import io
import zipfile
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from extractor import format_display_number

def apply_cell_shading(cell, fill_hex):
    """Applies background color fill to a Word table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    # Remove existing shading if present
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    shd_xml = f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'
    tcPr.append(parse_xml(shd_xml))

def annotate_docx(source_docx_bytes, discrepancies, file_role='doc1', file1_name='File 1', file2_name='File 2', display_standard='Vietnam'):
    """
    Annotates a Word document with:
    1. Background shading on mismatched cells (Red #FEE2E2) and extra rows (Yellow #FEF3C7).
    2. Native Word comments with specific audit findings formatted in unified display standard.
    
    Returns: Bytes of the annotated .docx file.
    """
    if not source_docx_bytes:
        return None
        
    doc = Document(io.BytesIO(source_docx_bytes))
    
    file_self_name = file1_name if file_role == 'doc1' else file2_name
    file_other_name = file2_name if file_role == 'doc1' else file1_name
    
    # Track cells already commented to avoid duplicate balloon clutter
    commented_cells = set()
    
    for disc in discrepancies:
        t_idx = disc.get('doc_table_idx', disc['table_idx'])
        r_idx = disc.get('doc_row_idx', disc['row_idx'])
        c_idx = disc.get('doc_col_idx', disc.get('col_idx', 1))
        disc_type = disc['type']
        label = disc.get('label', '')
        
        if t_idx - 1 >= len(doc.tables):
            continue
        table = doc.tables[t_idx - 1]
        
        if r_idx - 1 >= len(table.rows):
            continue
        row = table.rows[r_idx - 1]
        
        if disc_type == 'unmatched_table':
            # Highlight first row of the unmatched table in gentle amber
            if len(table.rows) > 0:
                first_row = table.rows[0]
                for cell in first_row.cells:
                    apply_cell_shading(cell, "FEF3C7")
                    
                cell_key = (t_idx, 1, 1)
                if cell_key not in commented_cells and len(first_row.cells) > 0:
                    commented_cells.add(cell_key)
                    target_cell = first_row.cells[0]
                    p = target_cell.paragraphs[0] if target_cell.paragraphs else target_cell.add_paragraph()
                    if not p.runs:
                        p.add_run(target_cell.text)
                    
                    rows_cnt = disc.get('rows', len(table.rows))
                    cols_cnt = disc.get('cols', len(table.columns))
                    comment_text = (
                        f"⚠️ [WordCompare] CẢNH BÁO CẤU TRÚC: BẢNG BỊ GỘP/TÁCH — CHƯA ĐỐI CHIẾU SỐ LIỆU\n"
                        f"• Tình trạng: Bảng {t_idx} ({rows_cnt} dòng × {cols_cnt} cột) không có bảng đối ứng 1:1 tương đương trong '{file_other_name}'.\n"
                        f"• Lý do: Có sự gộp hoặc tách bảng giữa 2 tài liệu (ví dụ: 2 bảng ở File 1 được gộp thành 1 bảng lớn ở File 2 hoặc ngược lại).\n"
                        f"• Trạng thái: CHƯA ĐỐI CHIẾU TỰ ĐỘNG số liệu bảng này (để tránh báo lỗi lệch dây chuyền cho các bảng phía sau).\n"
                        f"• Khuyến nghị: Vui lòng rà soát thủ công số liệu của bảng này với bảng tương đương trong '{file_other_name}'."
                    )
                    try:
                        doc.add_comment(runs=p.runs, text=comment_text, author="WordCompare Bot", initials="WC")
                    except Exception:
                        pass
        elif disc_type == 'extra_row_self':
            # Highlight all cells in the extra row
            for cell in row.cells:
                apply_cell_shading(cell, "FEF3C7")
                
            cell_key = (t_idx, r_idx, 1)
            if cell_key not in commented_cells and len(row.cells) > 0:
                commented_cells.add(cell_key)
                target_cell = row.cells[0]
                p = target_cell.paragraphs[0] if target_cell.paragraphs else target_cell.add_paragraph()
                if not p.runs:
                    p.add_run(target_cell.text or " ")
                elif not p.text.strip():
                    p.runs[0].text = " "
                
                if disc.get('is_blank'):
                    comment_text = (
                        f"ℹ️ [WordCompare] CẢNH BÁO DÒNG ĐỆM TRỐNG:\n"
                        f"• Tình trạng: Dòng đệm trống định dạng này chỉ xuất hiện ở '{file_self_name}', không có trong '{file_other_name}'.\n"
                        f"• Đánh giá: Không ảnh hưởng đến số liệu tài chính nhưng gây khác biệt số dòng bảng biểu.\n"
                        f"• Hệ thống đã tự động gióng đúng các dòng số liệu phía sau."
                    )
                else:
                    comment_text = (
                        f"ℹ️ [WordCompare] Dòng bổ sung (Thừa dòng):\n"
                        f"Dòng này có trong '{file_self_name}' nhưng không tìm thấy trong '{file_other_name}'.\n"
                        f"• Chỉ tiêu / Nhãn: {label}"
                    )
                try:
                    doc.add_comment(runs=p.runs, text=comment_text, author="WordCompare Bot", initials="WC")
                except Exception:
                    pass
        elif disc_type == 'merged_table_info':
            # Do NOT apply aggressive red/yellow shading! Keep cell formatting natural and clean.
            cell_key = (t_idx, r_idx, 'merged_info')
            if cell_key not in commented_cells and len(row.cells) > 0:
                commented_cells.add(cell_key)
                target_cell = row.cells[0]
                p = target_cell.paragraphs[0] if target_cell.paragraphs else target_cell.add_paragraph()
                if not p.runs:
                    p.add_run(target_cell.text or " ")
                elif not p.text.strip():
                    p.runs[0].text = " "
                
                info_type = disc.get('info_type', '')
                if info_type == 'split_source':
                    comment_text = (
                        f"ℹ️ [WordCompare] THÔNG TIN GỘP/TÁCH BẢNG (ĐÃ TỰ ĐỘNG ĐỐI CHIẾU):\n"
                        f"• Tình trạng: Bảng {t_idx} ({disc.get('self_rows', len(table.rows))} dòng) trong tài liệu này tương ứng với một phần của Bảng {disc.get('target_table', '?')} ({disc.get('other_rows', '?')} dòng) trong '{file_other_name}'.\n"
                        f"• Ghép cùng: Bảng {disc.get('partner_table', '?')} của tài liệu này.\n"
                        f"• Kết quả: Hệ thống đã tự động dựng lại cấu trúc và đối chiếu số liệu từng ô khớp 100%."
                    )
                elif info_type == 'merged_target_header':
                    src_tables = disc.get('source_tables', [])
                    src_rows = disc.get('source_rows', [])
                    s0 = f"Bảng {src_tables[0]} ({src_rows[0]} dòng)" if len(src_tables) > 0 else "Bảng trước"
                    s1 = f"Bảng {src_tables[1]} ({src_rows[1]} dòng)" if len(src_tables) > 1 else "Bảng sau"
                    comment_text = (
                        f"ℹ️ [WordCompare] THÔNG TIN GỘP/TÁCH BẢNG (ĐÃ TỰ ĐỘNG ĐỐI CHIẾU):\n"
                        f"• Tình trạng: Bảng {t_idx} ({disc.get('self_rows', len(table.rows))} dòng) trong tài liệu này được gộp từ {s0} và {s1} của '{file_other_name}'.\n"
                        f"• Kết quả: Hệ thống đã tự động phân tách cấu trúc và đối chiếu số liệu từng ô khớp 100%."
                    )
                elif info_type == 'merged_target_part2':
                    comment_text = (
                        f"ℹ️ [WordCompare] THÔNG TIN GỘP/TÁCH BẢNG (PHẦN TIẾP THEO):\n"
                        f"• Tình trạng: Từ dòng {disc.get('split_row', r_idx)} này trở đi tương ứng với Bảng {disc.get('source_table', '?')} ({disc.get('source_rows', '?')} dòng) của '{file_other_name}'.\n"
                        f"• Kết quả: Hệ thống đã tự động phân tách cấu trúc và đối chiếu số liệu từng ô khớp 100%."
                    )
                else:
                    comment_text = f"ℹ️ [WordCompare] THÔNG TIN GỘP/TÁCH BẢNG:\n• {disc.get('note', 'Bảng đã được tự động phân tách và đối chiếu khớp số liệu.')}"
                
                try:
                    doc.add_comment(runs=p.runs, text=comment_text, author="WordCompare Bot", initials="WC")
                except Exception:
                    pass
        else:
            if c_idx - 1 >= len(row.cells):
                continue
            cell = row.cells[c_idx - 1]
            apply_cell_shading(cell, "FEE2E2")
            
            cell_key = (t_idx, r_idx, c_idx)
            if cell_key not in commented_cells:
                commented_cells.add(cell_key)
                p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
                if not p.runs:
                    p.add_run(cell.text)
                
                if disc_type == 'num_mismatch':
                    val_self = disc['val_self']
                    val_other = disc['val_other']
                    diff = disc['diff']
                    s_self = format_display_number(val_self, display_standard)
                    s_other = format_display_number(val_other, display_standard)
                    s_diff = format_display_number(diff, display_standard, show_sign=True)
                    comment_text = (
                        f"⚠️ [WordCompare] Sai lệch số liệu:\n"
                        f"• {file_self_name}: {s_self}\n"
                        f"• {file_other_name}: {s_other}\n"
                        f"• Chênh lệch: {s_diff}\n"
                        f"• Chỉ tiêu: {label}"
                    )
                elif disc_type == 'format_mismatch':
                    comment_text = (
                        f"⚠️ [WordCompare] Lệch định dạng số / chữ:\n"
                        f"• {file_self_name}: {disc['val_self']}\n"
                        f"• {file_other_name}: {disc['val_other']}\n"
                        f"• Chỉ tiêu: {label}"
                    )
                else: # text mismatch
                    comment_text = (
                        f"⚠️ [WordCompare] Sai lệch nội dung văn bản:\n"
                        f"• {file_self_name}: {disc['val_self']}\n"
                        f"• {file_other_name}: {disc['val_other']}\n"
                        f"• Chỉ tiêu: {label}"
                    )
                
                try:
                    doc.add_comment(runs=p.runs, text=comment_text, author="WordCompare Bot", initials="WC")
                except Exception:
                    pass

    out_buf = io.BytesIO()
    doc.save(out_buf)
    return out_buf.getvalue()

def create_zip_package(doc1_annotated_bytes, doc1_name, doc2_annotated_bytes, doc2_name, excel_bytes):
    """
    Creates a ZIP archive containing:
    1. Annotated Doc 1 (.docx)
    2. Annotated Doc 2 (.docx)
    3. Formatted Comparison Report (.xlsx)
    """
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        if doc1_annotated_bytes:
            clean_name1 = doc1_name if doc1_name.endswith('.docx') else f"{doc1_name}.docx"
            zf.writestr(f"[Da_Comment]_{clean_name1}", doc1_annotated_bytes)
        if doc2_annotated_bytes:
            clean_name2 = doc2_name if doc2_name.endswith('.docx') else f"{doc2_name}.docx"
            zf.writestr(f"[Da_Comment]_{clean_name2}", doc2_annotated_bytes)
        if excel_bytes:
            zf.writestr("Bao_cao_doi_chieu.xlsx", excel_bytes)
            
    return zip_buffer.getvalue()
