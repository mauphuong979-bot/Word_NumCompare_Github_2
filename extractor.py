import re
import pandas as pd
from docx import Document

def clean_data(text):
    """Removes non-printable characters and extra whitespace."""
    if not text:
        return ""
    cleaned = "".join(ch for ch in text if ord(ch) >= 32 or ch in "\n\r\t")
    return cleaned.strip()

def parse_number(text, format_type):
    """
    STRICTLY parses a string into a float based on format_type.
    Uses smart normalization (replace) to handle commas, dots, negatives, and percentages.
    """
    if not text:
        return None
    text = text.replace('\xa0', ' ').strip()
    is_negative = False
    inner_text = text
    
    # Strip enclosing parentheses or trailing/leading minus
    if inner_text.startswith('(') and inner_text.endswith(')'):
        is_negative = True
        inner_text = inner_text[1:-1].strip()
    elif inner_text.startswith('-'):
        is_negative = True
        inner_text = inner_text[1:].strip()
    elif inner_text.endswith('-'):
        is_negative = True
        inner_text = inner_text[:-1].strip()
        
    # Strip percentage if present
    if inner_text.endswith('%'):
        inner_text = inner_text[:-1].strip()

    if not inner_text or not re.match(r'^[0-9,.]+$', inner_text):
        return None
        
    try:
        if format_type == 'Vietnam':
            # Chuẩn Việt Nam: Dấu chấm (.) là hàng nghìn -> xóa bỏ; Dấu phẩy (,) là thập phân -> đổi thành chấm
            val_str = inner_text.replace('.', '').replace(',', '.')
        else:
            # Chuẩn US / Quốc tế: Dấu phẩy (,) là hàng nghìn -> xóa bỏ; Dấu chấm (.) là thập phân
            val_str = inner_text.replace(',', '')
                
        val = float(val_str)
        return -val if is_negative else val
    except (ValueError, TypeError):
        return None

def format_display_number(val, display_standard="Vietnam", show_sign=False):
    """
    Formats a numeric float/int into a unified standard:
    - 'Vietnam': Thousand separator = '.', Decimal separator = ',' (e.g. 1.234.567 or 1.234.567,89)
    - 'US': Thousand separator = ',', Decimal separator = '.' (e.g. 1,234,567 or 1,234,567.89)
    Intelligent decimal handling: if number is an exact integer, don't force meaningless decimals.
    If it has fractional parts, display up to 4 significant decimal places, stripping trailing zeroes.
    """
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return ""
    if not isinstance(val, (int, float)):
        return str(val)
        
    sign_str = "+" if (show_sign and val > 0) else ""
    is_neg = val < 0
    abs_val = abs(val)
    
    # Check if integer
    if isinstance(abs_val, int) or abs_val.is_integer():
        int_part = int(round(abs_val))
        if "Vietnam" in str(display_standard) or "Việt Nam" in str(display_standard):
            res = f"{int_part:,}".replace(',', '.')
        else:
            res = f"{int_part:,}"
    else:
        formatted_dec = f"{abs_val:,.4f}".rstrip('0').rstrip('.')
        if "Vietnam" in str(display_standard) or "Việt Nam" in str(display_standard):
            res = formatted_dec.replace(',', '#').replace('.', ',').replace('#', '.')
        else:
            res = formatted_dec
            
    prefix = "-" if is_neg else sign_str
    return f"{prefix}{res}"


def detect_number_format(doc_source):
    """
    Automatically detects whether a Word document uses 'Vietnam' or 'US' number format.
    Returns: {
        'format': 'Vietnam' | 'US',
        'confidence': float (0-100),
        'vn_count': int,
        'us_count': int,
        'samples': list of strings
    }
    """
    if hasattr(doc_source, 'tables'):
        doc = doc_source
    else:
        if hasattr(doc_source, 'seek'):
            doc_source.seek(0)
        doc = Document(doc_source)
        
    vn_count = 0
    us_count = 0
    samples_vn = []
    samples_us = []
    
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                raw = c.text.strip()
                if not raw:
                    continue
                # Match VN pattern: 1.234 or 1.234.567 or 123,45
                if re.search(r'\b\d{1,3}(\.\d{3})+(,\d+)?\b', raw) or re.search(r'\b\d+,\d{1,2}\b', raw):
                    vn_count += 1
                    if len(samples_vn) < 3 and raw not in samples_vn:
                        samples_vn.append(raw)
                # Match US pattern: 1,234 or 1,234,567 or 123.45
                elif re.search(r'\b\d{1,3}(,\d{3})+(\.\d+)?\b', raw) or re.search(r'\b\d+\.\d{1,2}\b', raw):
                    us_count += 1
                    if len(samples_us) < 3 and raw not in samples_us:
                        samples_us.append(raw)
                        
    total = vn_count + us_count
    if us_count > vn_count:
        detected = 'US'
        conf = (us_count / total * 100) if total else 100.0
        samples = samples_us
    else:
        detected = 'Vietnam'
        conf = (vn_count / total * 100) if total else 100.0
        samples = samples_vn
        
    return {
        'format': detected,
        'confidence': round(conf, 1),
        'vn_count': vn_count,
        'us_count': us_count,
        'samples': samples
    }


def get_row_label(cells):
    """
    Finds the most descriptive label for the row (e.g. item name in financial statements).
    Prioritizes cells containing text/letters (Vietnamese or Latin).
    """
    for c in cells:
        raw = c.get('raw_text', '')
        if re.search(r'[a-zA-Z\u00C0-\u024F\u1EA0-\u1EF9]', raw):
            return raw
    for c in cells:
        if c.get('raw_text', ''):
            return c.get('raw_text', '')
    return "(Không có tiêu đề)"

def get_row_signature(cells):
    """
    Generates a signature for smart row alignment:
    Numbers are generalized to '<NUM>' so value changes don't disrupt alignment.
    """
    parts = []
    for c in cells:
        raw = c.get('raw_text', '').strip()
        if not raw:
            parts.append("<EMPTY>")
        elif c.get('is_num', False):
            parts.append("<NUM>")
        else:
            norm = re.sub(r'\s+', ' ', raw).lower()
            parts.append(norm)
    return " | ".join(parts)

def get_document_tables_summary(doc_source):
    """
    Returns a quick summary list of all tables in a Word document:
    Table Index, Rows count, Columns count, and Title/Preview snippet.
    """
    if hasattr(doc_source, 'tables'):
        doc = doc_source
    else:
        if hasattr(doc_source, 'seek'):
            doc_source.seek(0)
        doc = Document(doc_source)
    summary = []
    
    for t_idx, table in enumerate(doc.tables, 1):
        num_rows = len(table.rows)
        max_cols = 0
        for r in table.rows:
            seen_tc = set()
            u_cnt = 0
            for c in r.cells:
                if hasattr(c, '_tc'):
                    if c._tc in seen_tc:
                        continue
                    seen_tc.add(c._tc)
                u_cnt += 1
            if u_cnt > max_cols:
                max_cols = u_cnt
        num_cols = max_cols if max_cols > 0 else max((len(r.cells) for r in table.rows), default=0)
        
        # Extract preview from first non-empty text
        preview = ""
        for r in table.rows[:3]:
            row_texts = [clean_data(c.text) for c in r.cells if clean_data(c.text)]
            if row_texts:
                preview = " | ".join(row_texts[:3])
                if len(preview) > 60:
                    preview = preview[:57] + "..."
                break
        if not preview:
            preview = "(Bảng trống)"
            
        summary.append({
            'table_index': t_idx,
            'doc_table_idx': t_idx,
            'rows': num_rows,
            'cols': num_cols,
            'preview': preview
        })
    return summary

def extract_structured_tables(doc_source, format_type):
    """
    Extracts all tables in a document with rich cell & row metadata:
    Returns dict: { t_idx: [ { 'row_idx': r, 'label': ..., 'sig': ..., 'cells': [...] } ] }
    """
    if hasattr(doc_source, 'tables'):
        doc = doc_source
    else:
        if hasattr(doc_source, 'seek'):
            doc_source.seek(0)
        doc = Document(doc_source)
    tables_data = {}
    
    for t_idx, table in enumerate(doc.tables, 1):
        rows_list = []
        for r_idx, row in enumerate(table.rows, 1):
            cells_list = []
            seen_tc = set()
            c_idx = 1
            for cell in row.cells:
                if hasattr(cell, '_tc'):
                    if cell._tc in seen_tc:
                        continue
                    seen_tc.add(cell._tc)
                raw = clean_data(cell.text)
                num = parse_number(raw, format_type)
                cells_list.append({
                    'col_idx': c_idx,
                    'doc_table_idx': t_idx,
                    'doc_row_idx': r_idx,
                    'doc_col_idx': c_idx,
                    'raw_text': raw,
                    'num_val': num,
                    'is_num': num is not None
                })
                c_idx += 1
            
            row_label = get_row_label(cells_list)
            row_sig = get_row_signature(cells_list)
            
            rows_list.append({
                'row_idx': r_idx,
                'doc_table_idx': t_idx,
                'doc_row_idx': r_idx,
                'label': row_label,
                'sig': row_sig,
                'cells': cells_list
            })
        tables_data[t_idx] = rows_list
        
    return tables_data

# Backward compatibility helper if needed
def extract_table_data(file_path, format_type, mode='Number'):
    """Legacy helper for backward compatibility."""
    tables_data = extract_structured_tables(file_path, format_type)
    flat_data = []
    for t_idx, rows in tables_data.items():
        for row in rows:
            for cell in row['cells']:
                if not cell['raw_text']:
                    continue
                if (mode == 'Number' and cell['is_num']) or (mode == 'Text' and not cell['is_num']):
                    flat_data.append({
                        'Table': t_idx,
                        'Address': f"Table {t_idx}_R{row['row_idx']}C{cell['col_idx']}",
                        'Value': cell['num_val'],
                        'Raw': cell['raw_text']
                    })
    return pd.DataFrame(flat_data)
